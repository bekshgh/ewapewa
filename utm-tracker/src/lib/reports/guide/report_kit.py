# ═══ REPORT KIT (one file) — exec(open("report_kit.py").read()) — never edit ═══
# ═══════════════════════════════════════════════════════════════════════════
# REPORT KIT · PART 1 — LAYOUT ENGINE
# Writes workbook blocks in exactly the layout the AIESEC reports site parses:
#   title + subtitle → "  HEADING  —  subtitle" bars → KPI cards / tables / callouts / definitions.
# NEVER write cells by hand in analytics sheets: use only the methods below.
# ═══════════════════════════════════════════════════════════════════════════
import re, math, warnings, datetime as dt
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.formatting.rule import ColorScaleRule
warnings.filterwarnings("ignore")

# ── palette (soft blue) ──────────────────────────────────────────────────
BG, PANEL, HEADER, HDR_LIGHT = "F4F7FC", "EAF0FA", "2D5FA6", "4A7BC2"
ACCENT = ["3A6BBF", "7EB3E8", "A8C8F0", "C5DEFA"]
GREEN, AMBER = "5B9E6F", "D4882A"
ROW_BASE, ROW_ALT, LINE = "FFFFFF", "F0F5FC", "D5E0F0"
TXT_DARK, TXT_MID, TXT_LIGHT = "1A2B4A", "4A6080", "8098B8"
FONT = "Calibri"

# ── number formats: the site reads meaning from the FORMAT, so always pass one ──
F_INT, F_1, F_2, F_3 = "#,##0", "0.0", "0.00", "0.000"
F_PCT, F_PCT0 = "0.0%", "0%"
F_SIGN, F_SIGN2 = "+0;-0;0", "+0.00;-0.00;0.00"
F_DATE = "yyyy-mm-dd"
F_PP = '+0.0" pp";-0.0" pp";0.0" pp"'
F_YRS = '+0.0" y";-0.0" y";0.0" y"'
F_DAYS = '0" days"'
DATE = "date"  # column type marker: real Excel date cell


def py(v):
    """numpy / pandas value → plain python (NaN → None). Excel must never see numpy types."""
    if v is None:
        return None
    if isinstance(v, (np.bool_, bool)):
        return bool(v)
    if isinstance(v, (np.integer,)):
        return int(v)
    if isinstance(v, (float, np.floating)):
        return None if (math.isnan(v) or math.isinf(v)) else float(v)
    if isinstance(v, pd.Timestamp):
        return None if pd.isna(v) else v.to_pydatetime()
    if v is pd.NaT:
        return None
    return v


def mean_or_none(s):
    s = pd.to_numeric(pd.Series(s), errors="coerce").dropna()
    return float(s.mean()) if len(s) else None


def share_true(s):
    """Share of True among non-null values (None when nothing is known)."""
    s = pd.Series(s).dropna()
    return float(s.astype(bool).mean()) if len(s) else None


class Sheet:
    """One analytics sheet. W = number of grid columns (B …). Every heading spans the full width."""

    def __init__(self, wb, name, W=20, tab=HEADER, first=False):
        ws = wb.active if first else wb.create_sheet(name)
        ws.title = name
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = tab
        ws.column_dimensions["A"].width = 2
        for c in range(2, W + 3):
            ws.column_dimensions[CL(c)].width = 10.5
        self.ws, self.W, self.r = ws, W, 2
        self.c1, self.c2 = 2, W + 1
        self.pos = {}  # section title → {hdr, first, last, col:{label:col_index}}

    # ── low level ─────────────────────────────────────────────────────────
    def _cell(self, r, c1, c2, v, bg=None, fg=TXT_DARK, sz=10, bold=False, italic=False,
              h="left", wrap=False, fmt=None, border=None):
        ws = self.ws
        if c2 > c1:
            ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
        c = ws.cell(r, c1, py(v))
        if isinstance(c.value, str) and c.value.startswith("="):
            c.data_type = "s"  # a caption like "= 5% of registrants" must stay text, never become a formula
        c.font = Font(name=FONT, size=sz, bold=bold, italic=italic, color=fg)
        c.alignment = Alignment(horizontal=h, vertical="center", wrap_text=wrap)
        if fmt:
            c.number_format = fmt
        for cc in range(c1, c2 + 1):
            if bg:
                ws.cell(r, cc).fill = PatternFill("solid", fgColor=bg)
            if border:
                ws.cell(r, cc).border = border
        return c

    def _h(self, r, h):
        self.ws.row_dimensions[r].height = h

    def gap(self, h=8):
        self._h(self.r, h)
        self.r += 1

    # ── title ─────────────────────────────────────────────────────────────
    def title(self, text, subtitle):
        """text = "<Event> — KPI Dashboard" (event name BEFORE the dash). subtitle carries period + n."""
        self._h(2, 34)
        self._h(3, 20)
        self._cell(2, self.c1, self.c2, text, HEADER, "FFFFFF", 17, True)
        self._cell(3, self.c1, self.c2, subtitle, PANEL, TXT_MID, 10, italic=True)
        self.r = 5

    # ── heading ───────────────────────────────────────────────────────────
    def heading(self, title, sub=None):
        """Section heading. TWO leading spaces + UPPER-CASE title are what the parser looks for."""
        letters = re.sub(r"[^A-Za-zА-Яа-я]", "", title)
        assert sum(ch.isupper() for ch in letters) / max(len(letters), 1) >= 0.8 and "(" not in title, \
            f"heading must be UPPER-CASE (a trailing lower-case 's' / 'vs' is fine), no brackets: {title}"
        if self.r > 5:
            self.gap(10)
        text = "  " + title + (f"  —  {sub}" if sub else "")
        self._h(self.r, 22)
        self._cell(self.r, self.c1, self.c2, text, HDR_LIGHT, "FFFFFF", 10.5, True)
        self.r += 1

    # ── KPI cards ─────────────────────────────────────────────────────────
    def kpis(self, title, cards, sub=None):
        """cards: [(LABEL, value, number_format|None, caption)]. ≥2 cards. Rows of ≤6, balanced."""
        assert len(cards) >= 2, "a KPI block needs at least 2 cards"
        self.heading(title, sub)
        n = len(cards)
        rows = math.ceil(n / 6)
        base, extra = divmod(n, rows)
        sizes = [base + (1 if i < extra else 0) for i in range(rows)]
        i = 0
        for size in sizes:
            span = max(2, self.W // size)
            chunk = cards[i:i + size]
            i += size
            r0 = self.r
            self._h(r0, 18), self._h(r0 + 1, 34), self._h(r0 + 2, 18)
            for k, (label, value, fmt, caption) in enumerate(chunk):
                c1 = self.c1 + k * span
                c2 = c1 + span - 1
                top = Side(style="thin", color=LINE)
                self._cell(r0, c1, c2, label.upper(), ROW_BASE, TXT_LIGHT, 8, True, h="center",
                           border=Border(top=Side(style="medium", color=ACCENT[k % 4]), left=top, right=top))
                self._cell(r0 + 1, c1, c2, value, ROW_BASE, HEADER, 20, True, h="center", fmt=fmt,
                           border=Border(left=top, right=top))
                self._cell(r0 + 2, c1, c2, caption, ROW_BASE, TXT_MID, 8, italic=True, h="center",
                           border=Border(left=top, right=top, bottom=top))
            self.r = r0 + 3
            self.gap(8)

    # ── table ─────────────────────────────────────────────────────────────
    def table(self, title, cols, rows, sub=None, total=False, note=None, heat=None, diverging=False, heat_block=False):
        """cols: [(label, fmt)] — fmt None = text, DATE = Excel date, else a number format.
        rows: list of lists (None = empty). total=True → last row is the Total/All row.
        note: one line (>40 chars) printed under the table. heat=(i0,i1): colour scale over columns i0..i1,
        one scale per column (heat_block=True → one shared scale, for matrices like channel × day)."""
        assert rows, f"empty table: {title}"
        self.heading(title, sub)
        ws, n = self.ws, len(cols)
        spans = [1] * n
        if cols[0][1] is None:
            spans[0] = 3
        for i in range(1, n):
            if cols[i][1] is None:
                lens = [len(str(r[i])) for r in rows if r[i] is not None]
                if lens and max(lens) > 22:
                    spans[i] = 3
                elif lens and max(lens) > 12:
                    spans[i] = 2
        while sum(spans) > self.W:
            j = max(range(n), key=lambda k: (spans[k], k))
            assert spans[j] > 1, f"table too wide for the sheet ({title}): raise W"
            spans[j] -= 1
        starts = [self.c1 + sum(spans[:i]) for i in range(n)]
        col_of = {}
        for i, (lab, _) in enumerate(cols):
            col_of.setdefault(lab, starts[i])
        # header
        hl = max(math.ceil(len(lab) / (10 * spans[i])) for i, (lab, _) in enumerate(cols))
        self._h(self.r, 14 * hl + 10)
        hdr = self.r
        for i, (lab, _) in enumerate(cols):
            self._cell(hdr, starts[i], starts[i] + spans[i] - 1, lab, HEADER, "FFFFFF", 9, True, h="center", wrap=True)
        self.r += 1
        first = self.r
        for ri, row in enumerate(rows):
            assert len(row) == n, f"row {ri} has {len(row)} values, expected {n} ({title})"
            is_tot = total and ri == len(rows) - 1
            bg = PANEL if is_tot else (ROW_BASE if ri % 2 == 0 else ROW_ALT)
            self._h(self.r, 18)
            for i, v in enumerate(row):
                fmt = cols[i][1]
                if isinstance(v, tuple):          # (value, own_format) for mixed-format "Metric | Value" tables
                    v, fmt = v
                is_date = fmt == DATE
                v = py(v)
                if isinstance(v, str):
                    assert not v.startswith("  "), f"cell text must not start with 2 spaces: {v!r}"
                    assert i == 0 or not re.fullmatch(r"[-+]?[\d.,]+\s*(%|yrs?|pp)?", v), f"number stored as text: {v!r} ({title})"
                self._cell(self.r, starts[i], starts[i] + spans[i] - 1, v, bg, TXT_DARK, 10, bold=is_tot,
                           h="left" if (i == 0 or fmt is None) else "center",
                           fmt=(F_DATE if is_date else fmt),
                           border=Border(bottom=Side(style="thin", color=LINE)))
            self.r += 1
        last = self.r - 1
        if heat:
            a, b = heat
            bottom = last - (1 if total else 0)
            if heat_block:
                ranges = [f"{CL(starts[a])}{first}:{CL(starts[b] + spans[b] - 1)}{bottom}"]
            else:
                ranges = [f"{CL(starts[i])}{first}:{CL(starts[i] + spans[i] - 1)}{bottom}" for i in range(a, b + 1)]
            for rng in ranges:
                if diverging:
                    ws.conditional_formatting.add(rng, ColorScaleRule(start_type="min", start_color="F4B183",
                                                  mid_type="num", mid_value=0, mid_color="FFFFFF", end_type="max", end_color="9DC3E6"))
                else:
                    ws.conditional_formatting.add(rng, ColorScaleRule(start_type="min", start_color="FFFFFF",
                                                  end_type="max", end_color="7EB3E8"))
        if note:
            assert len(note) > 40, "note must be longer than 40 characters or it is read as a table row"
            self._h(self.r, 30)
            self._cell(self.r, self.c1, self.c2, note, None, TXT_MID, 9, italic=True, wrap=True)
            self.r += 1
        self.pos.setdefault(title, {"hdr": hdr, "first": first, "last": last - (1 if total else 0), "col": col_of})

    # ── callouts ("▸ text" lines) ─────────────────────────────────────────
    def callouts(self, title, items, sub=None):
        """items: str | (label, text). Tone from a leading glyph in the text: ⚠ = warning, ✓ = good."""
        self.heading(title, sub)
        for k, it in enumerate(items):
            label, text = it if isinstance(it, tuple) else (None, it)
            glyph = "⚠" if text.lstrip().startswith("⚠") else ("✓" if text.lstrip().startswith("✓") else "▸")
            text = re.sub(r"^[⚠✓▸]\s*", "", text)
            bg = ROW_BASE if k % 2 == 0 else ROW_ALT
            width_chars = (self.W - (3 if label else 1)) * 11
            self._h(self.r, max(20, 15 * math.ceil(len(text) / width_chars) + 6))
            if label:
                self._cell(self.r, self.c1, self.c1 + 2, f"{glyph} {label}", bg, ACCENT[0], 10, True, wrap=True)
                self._cell(self.r, self.c1 + 3, self.c2, text, bg, TXT_DARK, 10, wrap=True)
            else:
                self._cell(self.r, self.c1, self.c1, glyph, bg, AMBER if glyph == "⚠" else ACCENT[0], 11, True, h="center")
                self._cell(self.r, self.c1 + 1, self.c2, text, bg, TXT_DARK, 10, wrap=True)
            self.r += 1

    # ── definitions (Notes sheet) ─────────────────────────────────────────
    def definitions(self, title, pairs):
        """pairs: [(term, definition)] — definitions must be longer than ~25 characters."""
        self.heading(title)
        for k, (term, text) in enumerate(pairs):
            if len(text) <= 25:  # the parser tells definitions from table rows by length — keep them descriptive
                text += " — see the dashboard sheets for the numbers."
            bg = ROW_BASE if k % 2 == 0 else ROW_ALT
            self._h(self.r, max(20, 15 * math.ceil(len(text) / ((self.W - 3) * 11)) + 6))
            self._cell(self.r, self.c1, self.c1 + 2, term, bg, ACCENT[0], 10, True, wrap=True)
            self._cell(self.r, self.c1 + 3, self.c2, text, bg, TXT_DARK, 10, wrap=True)
            self.r += 1


def write_data_sheet(wb, name, frame, tab="B0BDD0"):
    """Plain data sheet (header in row 1, no merges). The site never opens it; keep contact columns OUT."""
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = tab
    ws.freeze_panes = "A2"
    for j, c in enumerate(frame.columns, 1):
        cell = ws.cell(1, j, str(c))
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=9)
        cell.fill = PatternFill("solid", fgColor=HEADER)
        ws.column_dimensions[CL(j)].width = max(11, min(30, len(str(c)) + 3))
    for i, row in enumerate(frame.itertuples(index=False), 2):
        for j, v in enumerate(row, 1):
            v = py(v)
            if isinstance(v, bool):
                v = "Yes" if v else "No"
            c = ws.cell(i, j, v)
            if isinstance(v, str) and v.startswith("="):
                c.data_type = "s"
            if isinstance(v, dt.datetime):
                c.number_format = F_DATE if (v.hour, v.minute, v.second) == (0, 0, 0) else "yyyy-mm-dd hh:mm"
    return ws


# ── helpers shared by both report types ─────────────────────────────────
def hhi(counts):
    s = pd.Series(counts, dtype=float)
    s = s / s.sum()
    return float((s ** 2).sum())


def concentration_label(h):
    return "Low" if h < 0.15 else ("Moderate" if h < 0.25 else "High")


def top_with_other(series, k, other_label="Other"):
    """value_counts, keep top-k, fold the rest into one 'Other …' row."""
    vc = series.dropna().astype(str).value_counts()
    if len(vc) <= k:
        return vc
    head, tail = vc.iloc[:k], vc.iloc[k:]
    return pd.concat([head, pd.Series({f"{other_label} ({len(tail)})": int(tail.sum())})])


def daterange_label(d0, d1):
    f = lambda d: f"{d.day} {d.strftime('%b')}"
    return f"{f(d0)} – {f(d1)}"

# ═══════════════════════════════════════════════════════════════════════════
# REPORT KIT · PART 2 — REGISTRATION REPORT
# build_registration(df, meta, path)  →  .xlsx the site reads with zero warnings.
#
# df = ONE ROW PER UNIQUE REGISTRANT (already de-duplicated), canonical columns:
#   ts            datetime   when the form was submitted (time of day if the export has it)
#   channel       str        tracked source, finest level (e.g. "TG Mailing #1", "Insta bio")
#   group         str        channel group (Telegram / Instagram / Email / Offline / Partners / Direct / Other)
#                            – if there is no finer split, channel = group
#   university    str        canonical short name (use norm_university)
#   study_level   str        canonical level (use norm_level): Bachelor Y1…Y4, Masters, Graduate, PhD, College, Foundation, Other
#   major         str        field of study, grouped to ≤ 12 categories
#   age           float      years (NaN allowed; drop impossible ages < 14 or > 60 to NaN)
#   wants_internship bool    interested in AIESEC internship (None if not asked)
#   new_to_aiesec bool       no previous AIESEC project (None if not asked)
# optional: prior_projects (str, ";"-separated, "" = none) · heard (str, self-reported source) ·
#           utm_key (str) · campaign (str)
# meta = dict(event="Project S'26", description="Project registrations", n_submissions=<raw rows>)
# NO names / e-mails / phones anywhere in df.
# ═══════════════════════════════════════════════════════════════════════════

UNI_RULES = [  # (keyword in lower-cased raw text, canonical name) — first match wins; extend when new values appear
    ("nazarbayev", "Nazarbayev University"), ("назарбаев", "Nazarbayev University"), (" nu ", "Nazarbayev University"),
    ("astana it", "Astana IT University"), ("aitu", "Astana IT University"), ("астана it", "Astana IT University"),
    ("narikbayev", "Maqsut Narikbayev Univ."), ("нарикбаев", "Maqsut Narikbayev Univ."), ("kazguu", "Maqsut Narikbayev Univ."),
    ("гумилев", "ENU (Gumilyov)"), ("гумилёв", "ENU (Gumilyov)"), ("enu", "ENU (Gumilyov)"),
    ("кбту", "KBTU"), ("kbtu", "KBTU"), ("kimep", "KIMEP"), ("sdu", "SDU"), ("satbayev", "Satbayev University"),
    ("coventry", "Coventry Univ. KZ"), ("международный университет астана", "Astana International Univ."),
    ("муит", "IITU"), ("iitu", "IITU"),
    ("зарубеж", "Foreign University"), ("foreign", "Foreign University"),
    ("не в астане", "Non-Astana University"), ("non-astana", "Non-Astana University"),
]


def norm_university(raw):
    if raw is None or (isinstance(raw, float) and math.isnan(raw)):
        return None
    t = f" {str(raw).lower()} "
    for key, name in UNI_RULES:
        if key in t:
            return name
    return "Other"


def norm_level(raw):
    """Free text → Bachelor Y1…Y4 / Masters / Graduate / PhD / College / Foundation / Other."""
    if raw is None or (isinstance(raw, float) and math.isnan(raw)):
        return None
    t = str(raw).lower()
    m = re.search(r"([1-4])\s*[-–]?\s*(курс|course|year|й курс)", t) or re.search(r"(?:year|y)\s*([1-4])\b", t)
    if m and ("бакалавр" in t or "bachelor" in t or "year" in t or "y" in t):
        return f"Bachelor Y{m.group(1)}"
    for key, name in [("магистр", "Masters"), ("master", "Masters"), ("phd", "PhD"), ("докторант", "PhD"),
                      ("выпускник", "Graduate"), ("graduate", "Graduate"), ("foundation", "Foundation"),
                      ("колледж", "College"), ("college", "College")]:
        if key in t:
            return name
    return "Other"


def level_stage(level):
    if level in ("Bachelor Y1", "Bachelor Y2", "Y1", "Y2"):
        return "Early (Y1–Y2)"
    if level in ("Bachelor Y3", "Bachelor Y4", "Y3", "Y4"):
        return "Senior (Y3–Y4)"
    if level in ("Graduate", "Masters", "PhD"):
        return "Grad / Postgrad"
    return "Pre-uni / Other"


STAGES = ["Early (Y1–Y2)", "Senior (Y3–Y4)", "Grad / Postgrad", "Pre-uni / Other"]
WEEKDAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]


def _prep(df):
    d = df.copy()
    for col in ["channel", "group", "university", "study_level", "major", "age", "wants_internship",
                "new_to_aiesec", "prior_projects", "heard", "utm_key", "campaign"]:
        if col not in d:
            d[col] = np.nan
    d["ts"] = pd.to_datetime(d["ts"])
    d["group"] = d["group"].fillna(d["channel"]).fillna("Direct / Other")
    d["channel"] = d["channel"].fillna(d["group"])
    d["utm_key"] = d["utm_key"].fillna(d["channel"].astype(str).str.lower())
    d["date"] = d["ts"].dt.normalize()
    d["stage"] = d["study_level"].map(level_stage)
    d["age"] = pd.to_numeric(d["age"], errors="coerce").where(lambda s: (s >= 14) & (s <= 60))
    for b in ("wants_internship", "new_to_aiesec"):
        d[b] = d[b].map(lambda v: None if v is None or (isinstance(v, float) and math.isnan(v)) else bool(v)).astype(object)
    return d


def _profile_row(g, top_uni):
    """Audience-quality columns shared by several tables."""
    return [
        share_true(g["university"].eq(top_uni).where(g["university"].notna())),
        share_true(g["stage"].eq("Early (Y1–Y2)").where(g["study_level"].notna())),
        share_true(g["wants_internship"]),
        share_true(g["new_to_aiesec"]),
    ]


def build_registration(df, meta, path, charts=True):
    d = _prep(df)
    N = len(d)
    ev = meta["event"]
    n_sub = int(meta.get("n_submissions", N))
    dups = max(0, n_sub - N)
    d0, d1 = d["date"].min(), d["date"].max()
    days = pd.date_range(d0, d1)
    D = len(days)
    has_time = bool((d["ts"].dt.hour + d["ts"].dt.minute).sum() > 0)
    top_uni = d["university"].value_counts().index[0] if d["university"].notna().any() else "Other"
    uni_short = top_uni if len(top_uni) <= 26 else top_uni[:25] + "…"
    ch_vc = d["channel"].value_counts()
    gr_vc = d["group"].value_counts()
    groups = list(gr_vc.index[:8])
    d["group8"] = d["group"].where(d["group"].isin(groups), "Other")
    if (d["group8"] == "Other").any() and "Other" not in groups:
        groups.append("Other")
    channels = list(ch_vc.index[:14])
    d["channel15"] = d["channel"].where(d["channel"].isin(channels), "Other channels")
    if (d["channel15"] == "Other channels").any():
        channels.append("Other channels")
    sub = f"{meta['description']} · {daterange_label(d0, d1)} · {N:,} unique registrants · {n_sub:,} form submissions"

    # ── per-day matrix ────────────────────────────────────────────────────
    per_day = d.groupby(["date", "group8"]).size().unstack(fill_value=0).reindex(days, fill_value=0).reindex(columns=groups, fill_value=0)
    tot = per_day.sum(axis=1)
    cum = tot.cumsum()
    z = (tot - tot.mean()) / (tot.std(ddof=0) if tot.std(ddof=0) else 1)
    peak_day = tot.idxmax()
    wd = pd.DataFrame({"tot": tot, "wd": [x.weekday() for x in tot.index]})
    wd_tab = wd.groupby("wd").agg(regs=("tot", "sum"), days=("tot", "size")).reindex(range(7), fill_value=0)
    wd_tab["avg"] = (wd_tab["regs"] / wd_tab["days"].replace(0, np.nan)).fillna(0)
    best_wd = int(wd_tab["avg"].idxmax())
    half_day = int((cum >= N / 2).values.argmax()) + 1

    wb = Workbook()

    # ═════════ 01 KPI Dashboard ═════════
    s = Sheet(wb, "01 · KPI Dashboard", 20, HEADER, first=True)
    s.title(f"{ev} — KPI Dashboard", sub)
    ch_share = ch_vc / N
    eff_ch = 1 / hhi(ch_vc)
    top_uni_share = float((d["university"] == top_uni).sum() / d["university"].notna().sum()) if d["university"].notna().any() else None
    s.kpis("HEADLINE KPIs", [
        ("Unique registrants", N, F_INT, f"after removing {dups} duplicates"),
        ("Form submissions", n_sub, F_INT, f"duplicate rate {dups / max(n_sub, 1):.1%}"),
        ("Campaign days", D, "0", f"{d0.day} {d0:%b} → {d1.day} {d1:%b}"),
        ("Avg regs / day", N / D, F_1, f"median day: {tot.median():.0f}"),
        ("Peak day regs", int(tot.max()), F_INT, f"{peak_day.day} {peak_day:%b} · {tot.max() / N:.0%} of all regs"),
        ("Active channels", len(ch_vc), "0", f"in {len(gr_vc)} channel groups"),
        ("Top channel share", float(ch_share.iloc[0]), F_PCT0, str(ch_share.index[0])),
        ("Top-3 channel share", float(ch_share.iloc[:3].sum()), F_PCT0, f"effective # channels: {eff_ch:.1f}"),
        ("Top group share", float(gr_vc.iloc[0] / N), F_PCT0, str(gr_vc.index[0])),
        ("Average age", mean_or_none(d["age"]), F_1, f"median {d['age'].median():.0f} · max {d['age'].max():.0f}" if d["age"].notna().any() else "age not collected"),
        ("Top university", top_uni_share, F_PCT0, f"share of registrants from {uni_short}"),
        ("New to AIESEC", share_true(d["new_to_aiesec"]), F_PCT0, "no previous AIESEC project"),
    ])
    sig = [
        f"{ch_share.index[0]} is the #1 channel → {ch_share.iloc[0]:.0%} of registrations ({int(ch_vc.iloc[0])} people)",
        f"Concentration: top-3 channels = {ch_share.iloc[:3].sum():.0%} · HHI {hhi(ch_vc):.2f} · effective # of channels = {eff_ch:.1f}",
        f"{gr_vc.index[0]} is the dominant group ({gr_vc.iloc[0] / N:.0%}) · top-2 groups = {gr_vc.iloc[:2].sum() / N:.0%}",
        f"Peak {peak_day.day} {peak_day:%b}: {int(tot.max())} regs = {tot.max() / max(tot.mean(), 1e-9):.1f}× the daily average",
    ]
    if top_uni_share is not None:
        sig.append(f"{uni_short} = {top_uni_share:.0%} of registrants; share by channel ranges from "
                   f"{min(_profile_row(g, top_uni)[0] or 0 for _, g in d.groupby('channel15')):.0%} to "
                   f"{max(_profile_row(g, top_uni)[0] or 0 for _, g in d.groupby('channel15')):.0%}")
    if dups:
        sig.append(f"Data hygiene: {dups} of {n_sub} submissions were duplicates (same person submitted twice) and were removed")
    s.callouts("KEY SIGNALS", sig)
    rows = []
    for rank, ch in enumerate(channels, 1):
        g = d[d["channel15"] == ch]
        rows.append([ch, g["group"].mode().iat[0], len(g), len(g) / N, rank, mean_or_none(g["age"])] + _profile_row(g, top_uni))
    rows.append(["All channels", None, N, 1.0, None, mean_or_none(d["age"])] + _profile_row(d, top_uni))
    s.table("CHANNEL HEAD-TO-HEAD", [("Channel", None), ("Group", None), ("Regs", F_INT), ("Share", F_PCT), ("Rank", "0"),
            ("Avg age", F_1), ("% Top univ.", F_PCT), ("% Year 1–2", F_PCT), ("% Want internship", F_PCT), ("% New to AIESEC", F_PCT)],
            rows, sub=f"volume + audience quality · top univ. = {top_uni}", total=True, heat=(5, 9))
    ws01 = s

    # ═════════ 02 Trend Dashboard ═════════
    Wt = max(20, 3 + len(groups) + 7)
    s = Sheet(wb, "02 · Trend Dashboard", Wt, "2E8B57")
    s.title(f"{ev} — Trend Dashboard", sub)
    hours = d["ts"].dt.hour.value_counts().reindex(range(24), fill_value=0) if has_time else None
    cards = [
        ("Peak day", f"{peak_day.day} {peak_day:%b}", None, f"{int(tot.max())} regs · {tot.max() / max(tot.mean(), 1e-9):.1f}× daily avg"),
        ("Avg regs / day", N / D, F_1, f"across {D} calendar days"),
        ("Days to 50% of volume", half_day, "0", f"of {D} days — front- vs back-loaded"),
        ("Strongest weekday", WEEKDAYS[best_wd], None, f"{wd_tab['avg'].max():.1f} regs per occurrence"),
    ]
    if has_time:
        ph = int(hours.idxmax())
        cards.append(("Peak hour", f"{ph:02d}:00", None, f"{hours.max() / N:.0%} of regs in that hour"))
    cards.append(("Spike days (z ≥ 1.5)", int((z >= 1.5).sum()), "0", f"{int((tot == 0).sum())} zero-registration days"))
    s.kpis("CAMPAIGN MOMENTUM SNAPSHOT", cards)
    dod = tot.diff().fillna(0)
    roll = tot.rolling(3, min_periods=1).mean()
    rows = []
    for i, day in enumerate(days):
        flag = "Spike" if z.iloc[i] >= 1.5 else ("Zero" if tot.iloc[i] == 0 else None)
        rows.append([day, f"{day.day} {day:%b}", WEEKDAYS[day.weekday()]] + [int(per_day.iloc[i][g]) for g in groups] +
                    [int(tot.iloc[i]), int(cum.iloc[i]), float(cum.iloc[i] / N), float(dod.iloc[i]), float(roll.iloc[i]), float(z.iloc[i]), flag])
    s.table("DAY-BY-DAY TABLE", [("Date", DATE), ("Day", None), ("Wkday", None)] + [(g, F_INT) for g in groups] +
            [("Total", F_INT), ("Cumul.", F_INT), ("Cumul. %", F_PCT0), ("DoD Δ", F_SIGN), ("3-day avg", F_1), ("z-score", F_1), ("Flag", None)],
            rows, sub="group split · cumulative · momentum · spike detection (z ≥ 1.5)")
    s.table("WEEKDAY PROFILE", [("Weekday", None), ("Regs", F_INT), ("Days in window", "0"), ("Avg / day", F_1), ("Share", F_PCT)],
            [[WEEKDAYS[i], int(wd_tab.loc[i, "regs"]), int(wd_tab.loc[i, "days"]), float(wd_tab.loc[i, "avg"]), float(wd_tab.loc[i, "regs"] / N)] for i in range(7)])
    if has_time:
        s.table("HOUR-OF-DAY PROFILE", [("Hour", None), ("Regs", F_INT), ("Share", F_PCT)],
                [[f"{h:02d}:00", int(hours[h]), float(hours[h] / N)] for h in range(24)], sub="local time of the form submission")
    ws02 = s

    # ═════════ 03 Channel Dashboard ═════════
    s = Sheet(wb, "03 · Channel Dashboard", 20, "1B7A8C")
    s.title(f"{ev} — Channel Dashboard", sub)
    rows = []
    for g_name in groups:
        g = d[d["group8"] == g_name]
        rows.append([g_name, int(g["channel"].nunique()), len(g), len(g) / N, mean_or_none(g["age"])] + _profile_row(g, top_uni))
    rows.append(["Total", int(d["channel"].nunique()), N, 1.0, mean_or_none(d["age"])] + _profile_row(d, top_uni))
    s.table("CHANNEL GROUP SUMMARY", [("Group", None), ("Channels", "0"), ("Regs", F_INT), ("Share", F_PCT), ("Avg age", F_1),
            ("% Top univ.", F_PCT), ("% Year 1–2", F_PCT), ("% Want internship", F_PCT), ("% New to AIESEC", F_PCT)],
            rows, total=True, heat=(4, 8))
    rows = []
    for g_name in groups:
        g = d[d["group8"] == g_name]
        rows.append([g_name, len(g)] + [float((g["stage"] == st).sum() / max(g["study_level"].notna().sum(), 1)) for st in STAGES])
    rows.append(["All", N] + [float((d["stage"] == st).sum() / max(d["study_level"].notna().sum(), 1)) for st in STAGES])
    s.table("AUDIENCE MIX BY GROUP", [("Group", None), ("n", F_INT)] + [(st, F_PCT) for st in STAGES], rows,
            sub="STUDY STAGE (row %)", total=True, heat=(2, 5))
    top5 = [u for u in d["university"].value_counts().index if u != "Other"][:5]
    d["uni6"] = d["university"].where(d["university"].isin(top5), "Other")
    rows = []
    for ch in channels:
        g = d[d["channel15"] == ch]
        k = max(g["university"].notna().sum(), 1)
        rows.append([ch, len(g)] + [float((g["uni6"] == u).sum() / k) for u in top5 + ["Other"]])
    k = max(d["university"].notna().sum(), 1)
    rows.append(["All channels", N] + [float((d["uni6"] == u).sum() / k) for u in top5 + ["Other"]])
    s.table("CHANNEL × UNIVERSITY", [("Channel", None), ("n", F_INT)] + [(u, F_PCT) for u in top5 + ["Other"]], rows,
            sub="top 5 (row %, heat = concentration)", total=True, heat=(2, 2 + len(top5)))
    base = _profile_row(d, top_uni)
    base_age = mean_or_none(d["age"])
    rows = []
    for ch in channels:
        g = d[d["channel15"] == ch]
        pr = _profile_row(g, top_uni)
        ga = mean_or_none(g["age"])
        rows.append([ch, len(g)] + [None if (a is None or b is None) else (a - b) * 100 for a, b in zip(pr, base)] +
                    [None if (ga is None or base_age is None) else ga - base_age])
    s.table("PERFORMANCE GAP vs CAMPAIGN AVERAGE", [("Channel", None), ("n", F_INT), ("Δ % Top univ.", F_PP), ("Δ % Year 1–2", F_PP),
            ("Δ % Want internship", F_PP), ("Δ % New to AIESEC", F_PP), ("Δ Avg age", F_YRS)], rows,
            sub="percentage points; blue = above avg, orange = below", heat=(2, 6), diverging=True)
    ws03 = s

    # ═════════ 04 Audience Dashboard ═════════
    s = Sheet(wb, "04 · Audience Dashboard", 20, "5B9FD4")
    s.title(f"{ev} — Audience Dashboard", sub)
    major_vc = d["major"].dropna().astype(str).value_counts()
    s.kpis("WHO REGISTERED", [
        ("Average age", mean_or_none(d["age"]), F_1, f"IQR {d['age'].quantile(.25):.0f}–{d['age'].quantile(.75):.0f} yrs" if d["age"].notna().any() else "age not collected"),
        ("Median age", float(d["age"].median()) if d["age"].notna().any() else None, "0", f"youngest {d['age'].min():.0f} · oldest {d['age'].max():.0f}" if d["age"].notna().any() else ""),
        ("Distinct universities", int(d["university"].nunique()), "0", "incl. Other / non-listed"),
        ("Top university share", top_uni_share, F_PCT0, top_uni),
        ("Top field share", float(major_vc.iloc[0] / major_vc.sum()) if len(major_vc) else None, F_PCT0, str(major_vc.index[0]) if len(major_vc) else ""),
        ("Want an internship", share_true(d["wants_internship"]), F_PCT0, f"new to AIESEC {share_true(d['new_to_aiesec']) or 0:.0%}"),
    ])

    def ranked(title, label, series, sub=None, share_label="Share", order=None, limit=15):
        vc = series.dropna().astype(str).value_counts()
        if order:
            vc = vc.reindex([o for o in order if o in vc.index]).dropna().astype(int)
        elif len(vc) > limit:
            vc = top_with_other(series, limit)
        if vc.empty:
            return
        k = int(vc.sum())
        cum_ = (vc / k).cumsum()
        s.table(title, [("#", "0"), (label, None), ("Regs", F_INT), (share_label, F_PCT), ("Cumul.", F_PCT)],
                [[i + 1, str(n), int(c), float(c / k), float(cum_.iloc[i])] for i, (n, c) in enumerate(vc.items())], sub=sub)

    ranked("STUDY LEVEL", "Level", d["study_level"], order=["Bachelor Y1", "Bachelor Y2", "Bachelor Y3", "Bachelor Y4", "Masters", "Graduate", "PhD", "College", "Foundation", "Other"])
    ranked("UNIVERSITY", "University", d["university"])
    ranked("MAJOR", "Major", d["major"])
    if d["age"].notna().any():
        bins = [(-1, 17, "≤17"), (17, 19, "18–19"), (19, 21, "20–21"), (21, 24, "22–24"), (24, 99, "25+")]
        band = pd.Series({lab: int(((d["age"] > lo) & (d["age"] <= hi)).sum()) for lo, hi, lab in bins})
        band = band[band > 0].sort_values(ascending=False)
        k = int(band.sum())
        s.table("AGE BAND", [("#", "0"), ("Age band", None), ("Regs", F_INT), ("Share", F_PCT), ("Cumul.", F_PCT)],
                [[i + 1, n, int(c), float(c / k), float((band / k).cumsum().iloc[i])] for i, (n, c) in enumerate(band.items())])
    if d["heard"].notna().any():
        ranked("HOW THEY HEARD", "Source", d["heard"], sub="self-reported")
    if d["prior_projects"].notna().any():
        proj = d["prior_projects"].fillna("").map(lambda x: [p.strip() for p in str(x).split(";") if p.strip()] or ["None"])
        pv = proj.explode().value_counts()
        s.table("PAST AIESEC PROJECTS", [("#", "0"), ("Project", None), ("Regs", F_INT), ("% of regs", F_PCT)],
                [[i + 1, str(n), int(c), float(c / N)] for i, (n, c) in enumerate(pv.head(14).items())],
                sub="multi-select: shares can add up to more than 100%")
    ws04 = s

    # ═════════ 05 Patterns & Outliers ═════════
    weekly = D > 45
    if weekly:
        wk = ((days - d0).days // 7)
        n_cols = int(wk.max()) + 1
    else:
        n_cols = D
    s = Sheet(wb, "05 · Patterns & Outliers", max(20, 3 + n_cols + 1), "C8780A")
    s.title(f"{ev} — Patterns & Outliers", sub)
    mat = d.groupby(["channel15", "date"]).size().unstack(fill_value=0).reindex(columns=days, fill_value=0)
    if weekly:
        mat = mat.T.groupby(wk).sum().T
        hdr_cols = [(f"W{i + 1}", F_INT) for i in range(n_cols)]
    else:
        hdr_cols = [(day.to_pydatetime(), F_INT) for day in days]
    rows = [[ch] + [int(v) for v in mat.loc[ch].values] + [int(mat.loc[ch].sum())] for ch in channels]
    rows.append(["All channels"] + [int(v) for v in mat.sum().values] + [N])
    # date headers must be real dates → write the header via a dedicated path
    _heatmap(s, "CHANNEL × DAY HEATMAP", hdr_cols, rows, weekly)
    rows = []
    for ch in channels:
        g = d[d["channel15"] == ch]
        by = g.groupby("date").size()
        rows.append([ch, len(g), int(by.size), g["date"].min().to_pydatetime(), g["date"].max().to_pydatetime(), int(by.max()), float(by.max() / len(g)),
                     "Burst" if by.max() / len(g) >= 0.5 else ("Wave" if by.max() / len(g) >= 0.25 else "Steady"),
                     int((g["date"].max() - g["date"].min()).days + 1)])
    s.table("CHANNEL BURSTINESS", [("Channel", None), ("Regs", F_INT), ("Active days", "0"), ("First reg", DATE), ("Last reg", DATE),
            ("Best-day regs", F_INT), ("Best-day share", F_PCT), ("Profile", None), ("Active span", F_DAYS)], rows,
            sub="how much of each channel's volume landed on its single best day")
    if d["age"].notna().any():
        a = d["age"].dropna()
        q1, q3 = a.quantile(.25), a.quantile(.75)
        lo, hi = q1 - 1.5 * (q3 - q1), q3 + 1.5 * (q3 - q1)
        s.table("AGE OUTLIERS", [("Metric", None), ("Value", F_1)],
                [["Mean", a.mean()], ["Median", a.median()], ["Std dev", a.std()], ["Min", a.min()], ["Q1", q1], ["Q3", q3], ["IQR", q3 - q1],
                 ["Lower fence", lo], ["Upper fence", hi], ["Below lower fence", (int((a < lo).sum()), "0")], ["Above upper fence", (int((a > hi).sum()), "0")], ["Max", a.max()]],
                sub="IQR rule")
        cuts = [(-1, 16, "≤16")] + [(x - 1, x, str(x)) for x in range(17, 25)] + [(24, 29, "25–29"), (29, 99, "30+")]
        s.table("AGE HISTOGRAM", [("Age", None), ("Regs", F_INT)], [[lab, int(((a > lo_) & (a <= hi_)).sum())] for lo_, hi_, lab in cuts])
    conc = []
    for name, ser in [("Channels", d["channel"]), ("Channel groups", d["group"]), ("Universities", d["university"]), ("Majors", d["major"])]:
        vc = ser.dropna().value_counts()
        if len(vc):
            h = hhi(vc)
            conc.append([name, float(vc.iloc[0] / vc.sum()), float(vc.iloc[:3].sum() / vc.sum()), h, 1 / h, concentration_label(h)])
    s.table("CONCENTRATION", [("Dimension", None), ("Top-1 share", F_PCT), ("Top-3 share", F_PCT), ("HHI", F_3), ("Effective #", F_1), ("Concentration", None)],
            conc, sub="HHI: <0.15 low · 0.15–0.25 moderate · >0.25 high · effective # = 1 / HHI")
    if d["heard"].notna().any():
        src = list(d["heard"].dropna().value_counts().index[:6])
        rows = []
        for g_name in groups:
            g = d[(d["group8"] == g_name) & d["heard"].notna()]
            if len(g) == 0:
                continue
            conf = float(g["heard"].astype(str).str.lower().str.contains(re.escape(g_name.lower().split()[0])).mean())
            rows.append([g_name, len(g)] + [float((g["heard"] == x).mean()) for x in src] + [conf])
        if rows:
            s.table("ATTRIBUTION CHECK", [("UTM group", None), ("n", F_INT)] + [(x, F_PCT) for x in src] + [("Self-report confirms UTM", F_PCT)], rows,
                    sub="UTM channel group × self-reported source (row %)", heat=(2, 2 + len(src)))
    if d["wants_internship"].notna().any():
        rows = []
        for st in STAGES:
            g = d[(d["stage"] == st) & d["wants_internship"].notna()]
            if len(g):
                yes = float(g["wants_internship"].astype(bool).mean())
                rows.append([st, len(g), yes, 1 - yes])
        s.table("STUDY STAGE × INTERNSHIP INTEREST", [("Study stage", None), ("n", F_INT), ("Yes", F_PCT), ("No", F_PCT)], rows, sub="row %", heat=(2, 3))
    s.table("DUPLICATE SUBMISSIONS", [("Metric", None), ("Value", F_INT)],
            [["Raw form submissions", n_sub], ["Unique registrants", N], ["Duplicates removed", dups], ["Duplicate rate", (dups / max(n_sub, 1), F_PCT)]],
            sub="one registrant per person; earliest submission kept")
    ws05 = s

    # ═════════ 06 Detailed Tables ═════════
    s = Sheet(wb, "06 · Detailed Tables", max(20, 3 + len(groups) + 3), "7A8FA6")
    s.title(f"{ev} — Detailed Tables", sub)
    u = d.groupby(["utm_key", "campaign", "channel", "group"], dropna=False).size().reset_index(name="n").sort_values("n", ascending=False)
    s.table("UTM-LEVEL DETAIL", [("UTM key", None), ("Campaign", None), ("Channel", None), ("Group", None), ("Regs", F_INT), ("Share", F_PCT), ("Rank", "0")],
            [[r.utm_key, None if pd.isna(r.campaign) else str(r.campaign), r.channel, r.group, int(r.n), float(r.n / N), i + 1] for i, r in enumerate(u.itertuples())],
            sub="every tracked link")

    def cross(title, label, col):
        sub_d = d[d[col].notna()]
        if sub_d.empty:
            return
        ct = sub_d.groupby([col, "group8"]).size().unstack(fill_value=0).reindex(columns=groups, fill_value=0)
        ct = ct.loc[ct.sum(axis=1).sort_values(ascending=False).index].head(20)
        k = int(sub_d.shape[0])
        cu = (ct.sum(axis=1) / k).cumsum()
        rows = [[str(i), *[int(x) for x in ct.loc[i]], int(ct.loc[i].sum()), float(ct.loc[i].sum() / k), float(cu.loc[i])] for i in ct.index]
        rows.append(["Total", *[int(ct[g].sum()) for g in groups], int(ct.values.sum()), float(ct.values.sum() / k), None])
        s.table(title, [(label, None)] + [(g, F_INT) for g in groups] + [("Total", F_INT), ("Share", F_PCT), ("Cumul.", F_PCT)], rows,
                sub="counts · heat = where each segment comes from", total=True, heat=(1, len(groups)))

    cross("UNIVERSITY × CHANNEL GROUP", "University", "university")
    cross("MAJOR × CHANNEL GROUP", "Major", "major")
    cross("STUDY LEVEL × CHANNEL GROUP", "Level", "study_level")
    cross("SELF-REPORTED SOURCE × CHANNEL GROUP", "Source", "heard")
    ws06 = s

    # ═════════ 07 Notes ═════════
    s = Sheet(wb, "07 · Notes", 20, AMBER)
    s.title("Assumptions, definitions & how to use this workbook", f"{ev} · registration analysis")
    s.definitions("SCOPE & CLEANING", [
        ("Source", f"{meta.get('source', 'registration export')} ({n_sub:,} form submissions)."),
        ("De-duplication", f"One registrant per person (same e-mail / phone); the earliest submission is kept. {dups} duplicates removed."),
        ("Fake & test rows", "Rows that are obviously test or fake entries are excluded before any count."),
        ("Window", f"{daterange_label(d0, d1)} ({D} calendar days), dates in the form's local time."),
    ])
    s.definitions("CHANNEL ATTRIBUTION", [
        ("Channel", "The tracked source of the registration (UTM source / medium / content). Registrations without a tag are 'Direct / Other'."),
        ("Group", "Channels are grouped into Telegram, Instagram, Email, Offline, Partners and Direct / Other for the group-level tables."),
    ])
    s.definitions("METRIC DEFINITIONS", [
        ("Share", "Registrations of a channel or segment divided by all unique registrants."),
        ("HHI", "Herfindahl index = sum of squared shares. 1 means everything came from one source; effective # = 1 / HHI."),
        ("z-score", "Daily registrations minus the daily mean, divided by the standard deviation; z ≥ 1.5 is flagged as a spike."),
        ("Study stage", "Early = Bachelor Y1–Y2; Senior = Bachelor Y3–Y4; Grad / Postgrad = Graduate, Masters, PhD; everything else = Pre-uni / Other."),
        ("Percent columns", "Percent columns use only registrants who answered that question; empty answers are not counted as 'No'."),
    ])
    ws07 = s

    # ═════════ 08 Charts (native Excel charts — the site ignores this sheet, humans like it) ═════════
    if charts:
        _registration_charts(wb, ws01, ws02, ws03, ws04, groups, D)

    # ═════════ 09 Cleaned Data (analytics columns only — NO names / e-mails / phones) ═════════
    keep = ["date", "ts", "channel", "group", "utm_key", "campaign", "university", "study_level", "major", "age",
            "wants_internship", "new_to_aiesec", "prior_projects", "heard"]
    write_data_sheet(wb, "09 · Cleaned Data", d[keep].rename(columns={"ts": "submitted_at"}))
    wb.save(path)
    return path


def _heatmap(s, title, hdr_cols, rows, weekly):
    """CHANNEL × DAY table: header cells are REAL dates (format 'd mmm') so the site shows '10 Aug'."""
    cols = [("Channel", None)] + [(f"c{i}", F_INT) for i in range(len(hdr_cols))] + [("Total", F_INT)]
    s.table(title, cols, rows, sub="bursts (mailings) vs always-on channels" if not weekly else "weekly buckets (campaign longer than 45 days)",
            total=True, heat=(1, len(hdr_cols)), heat_block=True)
    p = s.pos[title]
    for i, (lab, _) in enumerate(hdr_cols):
        c = s.ws.cell(p["hdr"], 2 + 3 + i)
        c.value = lab
        if not isinstance(lab, str):
            c.number_format = "d mmm"


def _registration_charts(wb, ws01, ws02, ws03, ws04, groups, D):
    from openpyxl.chart import BarChart, LineChart, Reference
    ch = wb.create_sheet("08 · Charts")
    ch.sheet_view.showGridLines = False
    ch.sheet_properties.tabColor = ACCENT[1]
    ch["B2"] = "Charts (built from the dashboard tables · not read by the site)"
    ch["B2"].font = Font(name=FONT, size=14, bold=True, color=HEADER)
    p = ws02.pos["DAY-BY-DAY TABLE"]
    ws = ws02.ws
    c_first = p["col"][groups[0]]
    bar = BarChart(); bar.type = "col"; bar.grouping = "stacked"; bar.overlap = 100
    bar.title = "Daily registrations by channel group"; bar.width, bar.height = 26, 9
    bar.add_data(Reference(ws, min_col=c_first, max_col=c_first + len(groups) - 1, min_row=p["hdr"], max_row=p["last"]), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=p["col"]["Day"], min_row=p["first"], max_row=p["last"]))
    ch.add_chart(bar, "B4")
    ln = LineChart(); ln.title = "Cumulative registrations"; ln.width, ln.height = 26, 8; ln.legend = None
    ln.add_data(Reference(ws, min_col=p["col"]["Cumul."], min_row=p["hdr"], max_row=p["last"]), titles_from_data=True)
    ln.set_categories(Reference(ws, min_col=p["col"]["Day"], min_row=p["first"], max_row=p["last"]))
    ch.add_chart(ln, "B24")
    p = ws01.pos["CHANNEL HEAD-TO-HEAD"]
    b2 = BarChart(); b2.type = "bar"; b2.title = "Registrations by channel"; b2.width, b2.height = 26, 10; b2.legend = None
    b2.add_data(Reference(ws01.ws, min_col=p["col"]["Regs"], min_row=p["hdr"], max_row=p["last"]), titles_from_data=True)
    b2.set_categories(Reference(ws01.ws, min_col=p["col"]["Channel"], min_row=p["first"], max_row=p["last"]))
    ch.add_chart(b2, "B41")
    p = ws04.pos["UNIVERSITY"]
    b3 = BarChart(); b3.type = "bar"; b3.title = "Registrations by university"; b3.width, b3.height = 26, 10; b3.legend = None
    b3.add_data(Reference(ws04.ws, min_col=p["col"]["Regs"], min_row=p["hdr"], max_row=p["last"]), titles_from_data=True)
    b3.set_categories(Reference(ws04.ws, min_col=p["col"]["University"], min_row=p["first"], max_row=p["last"]))
    ch.add_chart(b3, "B62")

# ═══════════════════════════════════════════════════════════════════════════
# REPORT KIT · PART 3 — FEEDBACK REPORT
# build_feedback(df, scales, themes, meta, path)  →  .xlsx the site reads with zero warnings.
#
# df = ONE ROW PER UNIQUE RESPONDENT (de-duplicated), canonical columns:
#   ts            datetime   when the response was submitted
#   <scale keys>  number     one numeric column per rating question (named in `scales`)
# optional segment columns (each makes one Segment Analysis table; skip what the form did not ask):
#   university (use norm_university) · study_stage (use level_stage(norm_level(..))) ·
#   channel_group (how they registered) · aiesec_bg ("New to AIESEC" / "Has AIESEC experience")
# optional: comment (str, ≤140 chars, a SHORT summary of low-scorers' open answers — no names)
#
# scales = [dict(key="org", label="Overall organization", short="Org", max=10), ...]
#          the FIRST scale with max=10 is the primary one (drives NPS); keep ≤ 8 scales.
# themes = dict(liked=[(theme, mentions, example_or_None)], pain=[...], requests=[...],
#               networking=[(label, mentions)], speakers=[(label, mentions)])  ← hand-coded from open answers;
#          every key optional; examples are SHORT verbatim quotes with no names / contacts.
# meta   = dict(event="Project S'26", description="Post-event participant survey",
#               n_registrants=<int or None>, source="<file name>", insights=[str,…], patterns=[(label, text),…])
# ═══════════════════════════════════════════════════════════════════════════

BANDS = {
    10: [(1, 3, "1–3 (Low)"), (4, 6, "4–6 (Mid)"), (7, 8, "7–8 (Good)"), (9, 10, "9–10 (Exc)")],
    5: [(1, 2, "1–2"), (3, 3, "3"), (4, 4, "4"), (5, 5, "5")],
}


def _band_def(mx):
    return BANDS.get(mx) or [(i, i, str(i)) for i in range(1, mx + 1)]


def build_feedback(df, scales, themes, meta, path, charts=True):
    d = df.copy()
    d["ts"] = pd.to_datetime(d["ts"])
    d["date"] = d["ts"].dt.normalize()
    n = len(d)
    ev = meta["event"]
    prim = next((sc for sc in scales if sc["max"] == 10), scales[0])
    keys = [sc["key"] for sc in scales]
    for sc in scales:
        d[sc["key"]] = pd.to_numeric(d[sc["key"]], errors="coerce")
        sc["_x"] = d[sc["key"]].dropna()
        sc["_mean"] = float(sc["_x"].mean())
        sc["_s10"] = sc["_mean"] / sc["max"] * 10
        sc["_top"] = float((sc["_x"] >= 0.9 * sc["max"]).mean())
        sc["_low"] = float((sc["_x"] <= 0.6 * sc["max"]).mean())
        sc["_fmt"] = F_2 if sc["max"] <= 5 else F_1
    d["_comp"] = pd.concat([d[sc["key"]] / sc["max"] * 10 for sc in scales], axis=1).mean(axis=1)
    px = prim["_x"]
    prom, det = int((px >= 9).sum()), int((px <= 6).sum())
    nps = (prom - det) / len(px) * 100 if prim["max"] == 10 else None
    d0, d1 = d["date"].min(), d["date"].max()
    ordered = sorted(scales, key=lambda sc: sc["_s10"])
    weakest, best = ordered[0], ordered[-1]
    seg_cols = [(c, h) for c, h in [("university", "UNIVERSITY"), ("study_stage", "STUDY STAGE"),
                                    ("channel_group", "REGISTRATION CHANNEL GROUP"), ("aiesec_bg", "AIESEC BACKGROUND")]
                if c in d and d[c].notna().any()]
    top_uni = None
    if "university" in d and d["university"].notna().any():
        vc = d["university"].value_counts()
        vc = vc[vc.index != "Other"]
        top_uni = vc.index[0] if len(vc) else None
    day1 = float((d["date"] == d0).mean())
    h48 = float(((d["ts"] - d["ts"].min()).dt.total_seconds() <= 48 * 3600).mean())
    sub = f"{meta['description']} · n = {n:,} respondents" + (
        f" · {n / meta['n_registrants']:.0%} of {meta['n_registrants']:,} registrants" if meta.get("n_registrants") else "") + \
        f" · responses {daterange_label(d0, d1)}"
    wb = Workbook()

    # ═════════ 01 KPI Dashboard ═════════
    s = Sheet(wb, "01 · KPI Dashboard", 20, "1A3A5C", first=True)
    s.title(f"{ev} — KPI Dashboard", sub)
    core = [(f"Avg {prim['label']}", prim["_mean"], F_1, f"out of {prim['max']} · median {px.median():g}")]
    if nps is not None:
        core.append((f"{prim['short']} NPS", nps, F_SIGN, f"{prom} promoters / {det} detractors"))
    core += [("Composite score", float(d["_comp"].mean()), F_1, f"all {len(scales)} ratings, normalised /10"),
             ("Total responses", n, F_INT, f"{n / meta['n_registrants']:.0%} of {meta['n_registrants']:,} registrants" if meta.get("n_registrants") else "usable responses"),
             ("Best-rated", best["label"], None, f"{best['_s10']:.1f}/10"),
             ("Weakest", weakest["label"], None, f"⚠ {weakest['_s10']:.1f}/10")]
    s.kpis("CORE EVENT METRICS", core)
    ops = [sc for sc in scales if sc is not prim]
    if len(ops) < 2:
        ops = scales
    s.kpis("OPERATIONAL SCORES", [(sc["label"], sc["_mean"], sc["_fmt"],
                                   ("⚠ lowest score in the survey" if sc is weakest else f"out of {sc['max']} · {sc['_top']:.0%} gave top ratings")) for sc in ops])
    who = [("Answered on day 1", day1, F_PCT0, f"{int((d['date'] == d0).sum())} responses on the first day"),
           ("Answered within 48h", h48, F_PCT0, "share of all responses")]
    for col, head in seg_cols:
        v = d[col].value_counts()
        who.append((f"Top {head.split(' ')[-1].lower() if col != 'channel_group' else 'channel'}", float(v.iloc[0] / v.sum()), F_PCT0, str(v.index[0])))
    s.kpis("WHO ANSWERED", who, sub="respondent profile")
    rows = []
    for sc in scales:
        row = [sc["label"], f"1–{sc['max']}", (sc["_mean"], sc["_fmt"]), (float(sc["_x"].median()), F_1), sc["_s10"], sc["_top"], sc["_low"]]
        if top_uni:
            a = d.loc[d["university"] == top_uni, sc["key"]].dropna()
            b = d.loc[d["university"] != top_uni, sc["key"]].dropna()
            row += [(a.mean() / sc["max"] * 10) if len(a) else None, (b.mean() / sc["max"] * 10) if len(b) else None,
                    ((a.mean() - b.mean()) / sc["max"] * 10) if len(a) and len(b) else None]
        rows.append(row)
    cols = [("Metric", None), ("Scale", None), ("Mean", F_2), ("Median", F_1), ("Score /10", F_1), ("Top-box %", F_PCT), ("Low %", F_PCT)]
    if top_uni:
        cols += [(f"{top_uni[:12]} /10", F_1), ("Others /10", F_1), ("Gap (top−others)", F_SIGN2)]
    s.table("SCORE SUMMARY", cols, rows,
            sub="ALL RESPONDENTS" + (f" vs {top_uni} / OTHERS" if top_uni else ""),
            note="Score /10 = mean ÷ scale maximum × 10 · Top-box = rating ≥ 90% of the scale · Low = rating ≤ 60% of the scale",
            heat=(len(cols) - 1, len(cols) - 1) if top_uni else None, diverging=True)
    sig = [f"Best-rated: {best['label']} {best['_s10']:.1f}/10 · weakest: {weakest['label']} {weakest['_s10']:.1f}/10 · gap {best['_s10'] - weakest['_s10']:.1f} pts"]
    if nps is not None:
        sig.append(f"{prim['short']} NPS {nps:.0f} · {prom} of {len(px)} rated 9–10 · {det} detractor(s)")
    if themes.get("liked"):
        t = max(themes["liked"], key=lambda x: x[1])
        sig.append(f"Most praised: {t[0]} ({t[1]} of {n})")
    if themes.get("pain"):
        t = max(themes["pain"], key=lambda x: x[1])
        sig.append(f"Biggest pain point: {t[0]} ({t[1]} of {n})")
    sig += meta.get("insights", [])
    s.callouts("KEY SIGNALS", sig)

    # ═════════ 02 Satisfaction & Themes ═════════
    s = Sheet(wb, "02 · Satisfaction & Themes", 20, "0E7C7B")
    s.title(f"{ev} — Satisfaction & Themes", f"What delighted respondents, what hurt, and what they ask for next time · hand-coded from {n} open answers")
    rows = [[sc["label"], (sc["_mean"], sc["_fmt"]), sc["_top"], 1 - sc["_top"] - sc["_low"], sc["_low"]] for sc in scales]
    s.table("RATING SCALES", [("Question", None), ("Mean", F_2), ("Top-box %", F_PCT), ("Mid %", F_PCT), ("Low %", F_PCT)], rows,
            sub="mean and share of answers by band", heat=(2, 4))
    for key, title, sub_ in [("liked", "WHAT THEY LIKED", "unprompted praise"), ("pain", "PAIN POINTS", "what went wrong"),
                             ("requests", "REQUESTS FOR NEXT TIME", None)]:
        if themes.get(key):
            lst = sorted(themes[key], key=lambda x: -x[1])
            s.table(title, [("Rank", "0"), ("Theme", None), ("Mentions", F_INT), ("% of resp.", F_PCT), ("Example (verbatim)", None)],
                    [[i + 1, t[0], t[1], t[1] / n, t[2] if len(t) > 2 else None] for i, t in enumerate(lst)], sub=sub_)

    # ═════════ 03 Operational Scorecard ═════════
    s = Sheet(wb, "03 · Operational Scorecard", 20, "E05C3A")
    s.title(f"{ev} — Operational Scorecard", "Where did execution succeed and where does it need improvement?")
    for mx in sorted({sc["max"] for sc in scales}, reverse=True):
        fam = [sc for sc in scales if sc["max"] == mx]
        bands = _band_def(mx)
        cols = [("Metric", None), ("Mean", F_2 if mx <= 5 else F_1), ("Med", F_1)]
        for _, _, lab in bands:
            cols += [(lab, F_INT), (f"{lab} %", F_PCT)]
        rows = []
        for sc in fam:
            x = sc["_x"]
            row = [sc["label"], sc["_mean"], float(x.median())]
            for lo, hi, _ in bands:
                c = int(((x >= lo) & (x <= hi)).sum())
                row += [c, c / len(x)]
            rows.append(row)
        s.table("SCORE DISTRIBUTION", cols, rows, sub=f"1–{mx} questions · count & share per rating band")
    if top_uni:
        rows = []
        for sc in scales:
            a = d.loc[d["university"] == top_uni, sc["key"]].dropna()
            b = d.loc[d["university"] != top_uni, sc["key"]].dropna()
            if len(a) and len(b):
                k = 10 / sc["max"]
                rows.append([sc["label"], a.mean() * k, a.median() * k, b.mean() * k, b.median() * k, (a.mean() - b.mean()) * k,
                             top_uni if a.mean() > b.mean() else ("Others" if a.mean() < b.mean() else "Tie")])
        if rows:
            s.table("OPERATIONAL SCORES BY UNIVERSITY", [("Metric", None), (f"{top_uni[:12]} mean", F_1), (f"{top_uni[:12]} med", F_1), ("Other mean", F_1),
                    ("Other med", F_1), ("Gap (top−other)", F_SIGN2), ("Winner", None)], rows,
                    sub=f"{top_uni} vs other universities · all scores normalised to /10",
                    note="Gap = difference of mean scores (normalised /10); positive = the top university rated higher than the others.",
                    heat=(5, 5), diverging=True)
    for key, title, sub_ in [("networking", "NETWORKING", "what actually worked (open answers)"), ("speakers", "SPEAKERS & TOPICS", "named as most useful")]:
        if themes.get(key):
            lst = sorted(themes[key], key=lambda x: -x[1])
            s.table(title, [("Rank", "0"), ("Format / speaker / topic", None), ("Mentions", F_INT), ("% of resp.", F_PCT)],
                    [[i + 1, t[0], t[1], t[1] / n] for i, t in enumerate(lst)], sub=sub_)

    # ═════════ 04 Segment Analysis ═════════
    s = Sheet(wb, "04 · Segment Analysis", 20, "6B52A0")
    s.title(f"{ev} — Segment Analysis", "How do respondent segments rate the event?")
    seg_scale_cols = [(f"{sc['short']} /{sc['max']}", sc["_fmt"]) for sc in scales]
    seg_cols_hdr = [("Segment", None), ("n", F_INT)] + seg_scale_cols + [("Composite", F_1)] + ([("Promoter %", F_PCT)] if nps is not None else [])

    def seg_row(label, g):
        row = [label, len(g)] + [mean_or_none(g[sc["key"]]) for sc in scales] + [mean_or_none(g["_comp"])]
        if nps is not None:
            row.append(float((g[prim["key"]] >= 9).sum() / max(g[prim["key"]].notna().sum(), 1)))
        return row

    def seg_table(title, series):
        vc = series.dropna().astype(str).value_counts()
        if vc.empty:
            return
        keep = list(vc.index[:8])
        ser = series.astype(str).where(series.notna())
        ser = ser.where(ser.isin(keep) | ser.isna(), "Other")
        labels = keep + (["Other"] if (ser == "Other").any() and "Other" not in keep else [])
        rows = [seg_row(k, d[ser == k]) for k in labels]
        rows.append(seg_row("All respondents", d))
        s.table(title, seg_cols_hdr, rows, sub="mean scores", total=True, heat=(2, len(seg_cols_hdr) - 1))

    for col, head in seg_cols:
        seg_table(head, d[col])
    dd = (d["date"] - d0).dt.days
    timing = pd.Series(np.select([dd == 0, dd <= 2, dd <= 6], ["Day 1 (event day)", "Days 2–3", "Days 4–7"], "Day 8+"), index=d.index)
    seg_table("RESPONSE TIMING", timing)

    # ═════════ 05 Patterns & Outliers ═════════
    s = Sheet(wb, "05 · Patterns & Outliers", 20, "D4882A")
    s.title(f"{ev} — Patterns & Outliers", "Where are the outliers? What patterns emerge across the data?")
    for mx in sorted({sc["max"] for sc in scales}, reverse=True):
        fam = [sc for sc in scales if sc["max"] == mx]
        cols = [("Score", "0")]
        for sc in fam:
            cols += [(sc["label"], F_INT), (f"{sc['label']} %", F_PCT)]
        rows = []
        for v in range(1, mx + 1):
            row = [v]
            for sc in fam:
                c = int((sc["_x"] == v).sum())
                row += [c, c / len(sc["_x"])]
            rows.append(row)
        s.table("SCORE FREQUENCY DISTRIBUTION", cols, rows, sub=f"1–{mx} scale · count (share)")
    low_mask = pd.concat([(d[sc["key"]] <= 0.6 * sc["max"]) for sc in scales], axis=1)
    d["_nlow"] = low_mask.sum(axis=1)
    out = d[d["_nlow"] > 0].sort_values(["_nlow", "_comp"], ascending=[False, True]).head(15)
    sub_out = f"any rating ≤ 60% of the scale ({int((d['_nlow'] > 0).sum())} respondents)"
    if out.empty:
        out = d.sort_values("_comp").head(5)
        sub_out = "no low ratings — the 5 lowest composite scores"
    has_comment = "comment" in d and d["comment"].notna().any()
    cols = [("Respondent", None)] + [(sc["short"], "0") for sc in scales] + [("# low", "0")] + ([("Key comment (summarised from open answers)", None)] if has_comment else [])
    rows = [[f"Respondent {i + 1}"] + [r[sc["key"]] for sc in scales] + [int(r["_nlow"])] + ([None if pd.isna(r["comment"]) else str(r["comment"])[:140]] if has_comment else [])
            for i, (_, r) in enumerate(out.iterrows())]
    s.table("OUTLIER RESPONDENTS", cols, rows, sub=sub_out)
    others = [sc for sc in scales if sc is not prim]
    if others:
        rows = []
        for sc in others:
            r_ = float(d[[prim["key"], sc["key"]]].dropna().corr().iloc[0, 1]) if len(d[[prim["key"], sc["key"]]].dropna()) > 3 else None
            rows.append([sc["label"], r_, None if r_ is None else ("Strong" if abs(r_) >= .6 else "Moderate" if abs(r_) >= .3 else "Weak")])
        s.table("RELATIONSHIPS", [("Metric", None), (f"Correl. with {prim['short']}", F_SIGN2), ("Strength", None)], rows,
                sub=f"what moves the {prim['label'].lower()} score?")
    g = d.groupby("date").agg(r=("ts", "size"), a=(prim["key"], "mean"), c=("_comp", "mean"))
    rows = [[day.to_pydatetime(), int(r.r), float(r.r / n), float(r.a), float(r.c)] for day, r in g.iterrows()]
    cum = 0
    for row in rows:
        cum += row[1]
        row[2] = cum / n
    s.table("RESPONSE TIMING", [("Date", DATE), ("Responses", F_INT), ("Cumul. %", F_PCT0), (f"Avg {prim['short'].lower()}", F_1), ("Avg composite", F_1)],
            rows, sub="post-event decay")
    cv = sorted(scales, key=lambda sc: -float(sc["_x"].std() / sc["_mean"]))[0]
    pats = [(f"{weakest['label']} is the weakest link", f"{weakest['label']} {weakest['_s10']:.1f}/10 vs composite {d['_comp'].mean():.1f} · ⚠ lowest of {len(scales)} ratings"),
            (f"{best['label']} is the strength", f"{best['label']} {best['_s10']:.1f}/10 · {best['_top']:.0%} gave top ratings"),
            ("Most volatile rating", f"{cv['label']}: std {cv['_x'].std():.2f} on a 1–{cv['max']} scale (CV {cv['_x'].std() / cv['_mean']:.0%})"),
            ("Ceiling effect", f"{prim['_top']:.0%} of {prim['label'].lower()} ratings are top-box — small differences between segments are noise")]
    if nps is not None:
        pats.append(("Detractors", f"{det} of {len(px)} respondents rated ≤ 6 on {prim['label'].lower()}"))
    pats += meta.get("patterns", [])
    s.callouts("KEY PATTERNS & ANOMALIES IDENTIFIED", pats)

    # ═════════ 06 Notes ═════════
    s = Sheet(wb, "06 · Notes", 20, "7A8FA6")
    s.title("Assumptions, definitions & how to use this workbook", f"{ev} · feedback analysis")
    s.definitions("DATA & CLEANING", [
        ("Source", f"{meta.get('source', 'feedback form export')} ({meta.get('n_rows', n)} rows)."),
        ("Removed rows", "Summary rows without a timestamp and repeat submissions (same e-mail; the latest one is kept) are removed before any count."),
        ("Sample", f"{n} unique respondents after de-duplication" + (f" ≈ {n / meta['n_registrants']:.0%} of registrants." if meta.get("n_registrants") else ".")),
    ])
    s.definitions("SCALES & METRICS", [
        ("1–10 questions", ", ".join(sc["label"] for sc in scales if sc["max"] == 10) or "No 1–10 questions in this form."),
        ("1–5 questions", ", ".join(sc["label"] for sc in scales if sc["max"] == 5) or "No 1–5 questions in this form."),
        ("Score /10", "Mean divided by the scale maximum and multiplied by 10, so 1–5 and 1–10 questions can be compared."),
        ("Top-box / Low", "Top-box = rating ≥ 90% of the scale maximum; Low = rating ≤ 60% of the scale maximum; Mid = everything in between."),
        ("NPS", f"Share of promoters (9–10) minus share of detractors (≤ 6) on '{prim['label']}', in points from −100 to +100." if nps is not None else "Not calculated: this form has no 1–10 question."),
        ("Composite", "Average of all ratings after normalising each to /10, computed per respondent and then averaged."),
    ])
    if any(themes.get(k) for k in ("liked", "pain", "requests")):
        s.definitions("OPEN-ANSWER THEMES", [
            ("Method", "Every open answer was read and hand-coded into themes; one answer can carry several themes, so mentions can add up to more than the number of respondents."),
            ("% of resp.", "Mentions divided by the number of respondents (not by the number of mentions)."),
        ])
    s.definitions("CAVEATS", [
        ("Sample", "Respondents are usually the most engaged attendees, so scores are likely higher than the average participant's experience."),
        ("Small segments", "Segment means with a small n move a lot with one answer — read them as direction, not as proof."),
    ])

    if charts:
        _feedback_charts(wb, scales, s)
    keep = ["date", "ts"] + keys + [c for c, _ in seg_cols] + (["comment"] if has_comment else [])
    write_data_sheet(wb, "08 · Cleaned Data", d[keep].rename(columns={"ts": "submitted_at"}))
    wb.save(path)
    return path


def _feedback_charts(wb, scales, notes_sheet):
    from openpyxl.chart import BarChart, Reference
    ch = wb.create_sheet("07 · Charts")
    ch.sheet_view.showGridLines = False
    ch["B2"] = "Charts (built from the dashboard tables · not read by the site)"
    ch["B2"].font = Font(name=FONT, size=14, bold=True, color=HEADER)
    ws = wb["01 · KPI Dashboard"]
    # locate the SCORE SUMMARY table by its header text
    hdr = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 2).value == "Metric")
    last = hdr
    while ws.cell(last + 1, 2).value and len(str(ws.cell(last + 1, 2).value)) < 60 and ws.cell(last + 1, 3).value is not None:
        last += 1
    col_s10 = next(c for c in range(2, 30) if ws.cell(hdr, c).value == "Score /10")
    b = BarChart(); b.type = "bar"; b.title = "Score /10 by metric"; b.width, b.height = 24, 10; b.legend = None
    b.add_data(Reference(ws, min_col=col_s10, min_row=hdr, max_row=last), titles_from_data=True)
    b.set_categories(Reference(ws, min_col=2, min_row=hdr + 1, max_row=last))
    ch.add_chart(b, "B4")

# ═══════════════════════════════════════════════════════════════════════════
# REPORT KIT · PART 4 — FORMAT CHECK (a Python mirror of the site's parser rules)
# check_report("File.xlsx")  →  prints PASS / FAIL and every problem the site would complain about.
# Run it on the finished file AND again after any manual edit. Fix every ✗ before delivering.
# ═══════════════════════════════════════════════════════════════════════════
import os, re, datetime as _dt
from openpyxl import load_workbook as _load

_SIG = {"registration": [r"Trend Dashboard", r"Channel Dashboard", r"Audience Dashboard"],
        "feedback": [r"Satisfaction & Themes", r"Operational Scorecard", r"Segment Analysis"]}
_SHEETS = {  # key: (sheet-name regex, [expected section regexes])
    "registration": {
        "overview": (r"KPI Dashboard", [r"HEADLINE KPIs", r"KEY SIGNALS", r"CHANNEL HEAD-TO-HEAD"]),
        "trend": (r"(Trend Dashboard|Daily Trend)", [r"MOMENTUM", r"DAY-BY-DAY", r"WEEKDAY", r"HOUR"]),
        "channels": (r"Channel Dashboard", [r"GROUP SUMMARY", r"AUDIENCE MIX", r"CHANNEL × UNIVERSITY", r"PERFORMANCE GAP"]),
        "audience": (r"(Audience Dashboard|Segment Analysis)", [r"WHO REGISTERED", r"STUDY LEVEL", r"UNIVERSITY", r"MAJOR", r"AGE BAND", r"HOW THEY HEARD", r"PAST AIESEC"]),
        "patterns": (r"Patterns? & Outliers", [r"HEATMAP", r"BURSTINESS", r"AGE OUTLIERS", r"CONCENTRATION", r"ATTRIBUTION"]),
        "tables": (r"Detailed Tables", [r"UTM-LEVEL"]),
        "methodology": (r"Notes", []),
    },
    "feedback": {
        "overview": (r"KPI Dashboard", [r"CORE EVENT METRICS", r"OPERATIONAL SCORES", r"WHO ANSWERED", r"SCORE SUMMARY", r"KEY SIGNALS"]),
        "themes": (r"Satisfaction & Themes", [r"SCALES", r"WHAT THEY LIKED", r"PAIN POINTS", r"REQUESTS"]),
        "scorecard": (r"Operational Scorecard", [r"SCORE DISTRIBUTION", r"BY UNIVERSITY", r"NETWORKING", r"SPEAKERS"]),
        "segments": (r"Segment Analysis", [r"UNIVERSITY", r"STUDY STAGE", r"CHANNEL", r"AIESEC", r"TIMING"]),
        "patterns": (r"Patterns? & Outliers", [r"FREQUENCY", r"OUTLIER", r"RELATIONSHIPS", r"RESPONSE TIMING", r"PATTERNS"]),
        "methodology": (r"Notes", []),
    },
}
_NEED_COLS = {  # (sheet key, section regex) → column-label regexes the site's charts look up
    "registration": [("overview", r"HEAD-TO-HEAD", [r"^Regs$", r"Avg age"]),
                     ("trend", r"DAY-BY-DAY", [r"^Day$", r"^Wkday$", r"^Total$", r"^Cumul\.?$"]),
                     ("trend", r"WEEKDAY", [r"Avg"]), ("trend", r"HOUR", [r"Regs"]),
                     ("channels", r"GROUP SUMMARY", [r"^Regs$"]),
                     ("audience", r"STUDY LEVEL|UNIVERSITY|MAJOR|AGE BAND|HOW THEY HEARD|PAST AIESEC", [r"^Regs$", r"Share|% of regs"]),
                     ("patterns", r"CONCENTRATION", [r"HHI", r"Concentration"]),
                     ("tables", r"UTM-LEVEL", [r"Regs"])],
    "feedback": [("overview", r"SCORE SUMMARY", [r"Score /10"]),
                 ("themes", r"LIKED|PAIN|REQUESTS", [r"Theme", r"Mentions", r"% of resp"]),
                 ("scorecard", r"NETWORKING|SPEAKERS", [r"Format|speaker|topic", r"Mentions", r"% of resp"]),
                 ("segments", r".*", [r"^n$"]),
                 ("patterns", r"RESPONSE TIMING", [r"Responses"]), ("patterns", r"RELATIONSHIPS", [r"Correl"])],
}
_DATE_TABLES = {"registration": [("trend", r"DAY-BY-DAY")], "feedback": [("patterns", r"RESPONSE TIMING")]}
_EMAIL = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
_GLYPHS = "▶▸►•⚠✓✔"


def _norm(name):
    return re.sub(r"\s+", " ", re.sub(r"^\s*\d+[a-z]?\s*[·.\-–]?\s*", "", name, flags=re.I)).strip()


def _upper_ratio(t):
    L = re.sub(r"[^A-Za-zА-Яа-яЁё]", "", t)
    return (len(re.sub(r"[^A-ZА-ЯЁ]", "", L)) / len(L)) if L else 0


def _heading(raw):
    if not isinstance(raw, str) or not re.match(r"^\s{2,}\S", raw):
        return None
    t = raw.strip()
    m = re.search(r"\s+[—–]\s+|\s+(?=\()|\s{2,}·\s{2,}", t)
    title = re.sub(r"\s+", " ", t[:m.start()] if m else t).strip()
    letters = re.sub(r"[^A-Za-zА-Яа-яЁё]", "", title)
    if len(letters) < 2 or _upper_ratio(title) < 0.7:
        return None
    return re.sub(r"^[A-Z]\d?\s*·\s*", "", title)


def check_report(path, kind=None, verbose=True):
    bad, warn = [], []
    E = lambda m: bad.append(m)
    W = lambda m: warn.append(m)
    if not path.lower().endswith(".xlsx"):
        E("file must be .xlsx (not .xlsm / .xls)")
    size = os.path.getsize(path)
    if size > 4 * 1024 * 1024:
        E(f"file is {size / 1e6:.1f} MB — the site accepts at most 4 MB (drop Raw Data sheets, keep Cleaned Data lean)")
    wb, wbv = _load(path), _load(path, data_only=True)
    names = {ws.title: _norm(ws.title) for ws in wb.worksheets}
    # type
    hits = {k: sum(any(re.fullmatch(p, n, re.I) for n in names.values()) for p in sig) for k, sig in _SIG.items()}
    full = [k for k, h in hits.items() if h == 3]
    if len(full) == 1:
        typ = full[0]
    else:
        part = [k for k, h in hits.items() if h * 2 > 3 and all(o == k or hits[o] == 0 for o in hits)]
        typ = part[0] if len(part) == 1 else None
    if typ is None:
        E("report type not recognised: a Registration file needs sheets 'Trend Dashboard', 'Channel Dashboard', 'Audience Dashboard'; "
          "a Feedback file needs 'Satisfaction & Themes', 'Operational Scorecard', 'Segment Analysis' (number prefix like '02 · ' is fine)")
        typ = kind or ("feedback" if "feedback" in os.path.basename(path).lower() else "registration")
    if kind and typ != kind:
        E(f"file looks like a {typ} report but {kind} was expected")
    formulas = nocache = 0
    first_title_done = False
    for key, (rx, expects) in _SHEETS[typ].items():
        ws_name = next((t for t, n in names.items() if re.fullmatch(rx, n, re.I)), None)
        if ws_name is None:
            E(f"sheet missing: '{rx}' (found: {list(names.values())})")
            continue
        ws, wsv = wb[ws_name], wbv[ws_name]
        merged = {(m.min_row, m.min_col): m for m in ws.merged_cells.ranges}
        covered = {(r, c) for m in ws.merged_cells.ranges for r in range(m.min_row, m.max_row + 1) for c in range(m.min_col, m.max_col + 1)
                   if (r, c) != (m.min_row, m.min_col)}
        # formulas / cached values / PII
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    formulas += 1
                    if wsv[c.coordinate].value is None:
                        nocache += 1
                if isinstance(c.value, str) and _EMAIL.search(c.value):
                    E(f"{ws_name}!{c.coordinate}: e-mail address in an analytics sheet (personal data)")
        heads = []
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and (c.row, c.column) not in covered:
                    h = _heading(c.value)
                    if h:
                        m = merged.get((c.row, c.column))
                        heads.append(dict(r=c.row, c=c.column, c2=m.max_col if m else c.column, title=h))
        lost = []
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and (c.row, c.column) not in covered and not _heading(c.value) and not c.value.startswith("="):
                    t0 = re.sub(r"\s+", " ", re.split(r"\s+[—–]\s+|\s+(?=\()", c.value.strip())[0])
                    mm = merged.get((c.row, c.column))
                    if len(re.sub(r"[^A-Za-z]", "", t0)) >= 4 and _upper_ratio(t0) >= 0.8 and (
                            (mm and mm.max_col - mm.min_col >= 2) or any(re.search(x, t0) for x in expects)):
                        lost.append((c.column, c.coordinate, t0))
        heads.sort(key=lambda x: (x["r"], x["c"]))
        for h in heads:
            if h["c2"] == h["c"]:
                h["c2"] = ws.max_column
        for rx_e in expects:
            if not any(re.search(rx_e, h["title"]) for h in heads):
                miss = next(((a, t0) for _, a, t0 in sorted(lost) if re.search(rx_e, t0)), None)
                if miss:
                    E(f"{ws_name}!{miss[0]}: heading '{miss[1]}' lost its 2 leading spaces — the site does not see it as a section (retype it with two spaces in front)")
                    continue
                W(f"{ws_name}: expected section '{rx_e}' not found (heading = 2 leading spaces + UPPER-CASE text, merged across the sheet)")
        if not heads:
            E(f"{ws_name}: no section headings found")
            continue
        # title / subtitle
        if not first_title_done:
            t, sub = ws.cell(2, 2).value, ws.cell(3, 2).value
            if not (isinstance(t, str) and " — " in t):
                E(f"{ws_name}!B2 must read '<Event name> — KPI Dashboard' (event name before the dash)")
            if not (isinstance(sub, str) and re.search(r"\d{1,2}\s+[A-Za-z]{3,}\s*[–—-]\s*\d{1,2}\s+[A-Za-z]{3,}", sub)):
                W(f"{ws_name}!B3 subtitle should contain the period like '10 Aug – 28 Aug' (site shows it as the report period)")
            if not (isinstance(sub, str) and re.search(r"\bn\s*=\s*[\d,]+|[\d,]+\s+unique registrants", sub, re.I)):
                W(f"{ws_name}!B3 subtitle should contain 'N unique registrants' (registration) or 'n = N respondents' (feedback)")
            first_title_done = True
        # blocks
        for hi, h in enumerate(heads):
            end = next((o["r"] - 1 for o in heads[hi + 1:] if o["c"] <= h["c2"] and o["c2"] >= h["c"]), ws.max_row)
            lines = []
            for r in range(h["r"] + 1, end + 1):
                items = []
                for c in range(h["c"], h["c2"] + 1):
                    if (r, c) in covered:
                        continue
                    cell = ws.cell(r, c)
                    if cell.value is None:
                        continue
                    m = merged.get((r, c))
                    items.append((c, min(m.max_col, h["c2"]) if m else c, cell, wsv.cell(r, c).value))
                lines.append((r, items))
            blocks, cur = [], []
            for r, items in lines:
                if items:
                    cur.append((r, items))
                elif cur:
                    blocks.append(cur)
                    cur = []
            if cur:
                blocks.append(cur)
            for blk in blocks:
                if h["c2"] < max(it[1] for _, items in blk for it in items):
                    E(f"{ws_name} / {h['title']}: content is wider than its heading — merge the heading across the whole table")
                _check_block(blk, ws_name, h["title"], key, typ, E, W)
        for skey, rx_t in _DATE_TABLES[typ]:
            if skey != key:
                continue
            for h in heads:
                if re.search(rx_t, h["title"]):
                    hdr_row = next((r for r in range(h["r"] + 1, h["r"] + 4) if ws.cell(r, h["c"]).value == "Date"), None)
                    if hdr_row is None:
                        E(f"{ws_name} / {h['title']}: first column must be labelled 'Date'")
                    else:
                        v = ws.cell(hdr_row + 1, h["c"]).value
                        if not isinstance(v, (_dt.datetime, _dt.date)):
                            E(f"{ws_name} / {h['title']}: dates must be real Excel dates (format yyyy-mm-dd), not text — the site takes the event date from here")
    if formulas >= 20 and nocache / max(formulas, 1) > 0.5:
        E(f"{nocache} of {formulas} formulas have no saved result — the site rejects the file. Use static values, or recalculate in LibreOffice/Excel and re-save")
    elif nocache:
        W(f"{nocache} formula cell(s) have no saved value and will show as empty")
    if verbose:
        print(f"{'FAIL' if bad else 'PASS'} · {typ} report · {os.path.basename(path)}")
        for m in bad:
            print("  ✗", m)
        for m in warn:
            print("  ! ", m)
        if not bad and not warn:
            print("  everything the site looks for is present")
    return dict(ok=not bad, type=typ, errors=bad, warnings=warn)


def _check_block(blk, sheet, title, key, typ, E, W):
    first = blk[0][1]
    txt = lambda it: it[3] if isinstance(it[3], str) else None
    up = lambda v: isinstance(v, str) and _upper_ratio(v) >= 0.8 and len(re.sub(r"[^A-Za-z]", "", v)) >= 2
    # KPI cards
    if len(blk) >= 2 and len(blk) % 3 != 1 and all(len(blk[i][1]) >= 2 and all(up(txt(it)) for it in blk[i][1]) for i in range(0, len(blk), 3)):
        return
    # callouts / definitions / prose
    if all(blk[i][1] and isinstance(blk[i][1][0][3], str) and blk[i][1][0][3].lstrip()[:1] in _GLYPHS and len(blk[i][1]) <= 2 for i in range(len(blk))):
        return
    if all(len(items) == 2 and all(isinstance(it[3], str) for it in items) for _, items in blk) and \
            sum(len(items[1][3]) for _, items in blk) / len(blk) > 25:
        return
    if len(blk) == 1 and len(first) == 1 and isinstance(first[0][3], str):
        return
    # table
    if len(first) < 2:
        W(f"{sheet} / {title}: block not recognised as KPI cards, table, callouts or definitions (header row needs ≥ 2 cells)")
        return
    hdr = [(it[0], it[1], it[3]) for it in first]
    rng = [(c1, c2 if c2 > c1 else (hdr[i + 1][0] - 1 if i + 1 < len(hdr) else c1)) for i, (c1, c2, _) in enumerate(hdr)]
    labels = [str(h[2]) if h[2] is not None else "" for h in hdr]
    for r, items in blk[1:]:
        for it in items:
            ci = next((i for i, (a, b) in enumerate(rng) if a <= it[0] <= b), None)
            if ci is None:
                W(f"{sheet} / {title}: a cell in row {r} lies outside every header column and will be dropped")
                continue
            v, cell = it[3], it[2]
            if isinstance(v, str) and re.fullmatch(r"\s*[-+]?[\d.,]+\s*(%|yrs?|pp|y)?\s*", v) and ci > 0:
                E(f"{sheet}!{cell.coordinate}: number stored as text ({v!r}) — write a real number with a number format (e.g. 0.0%)")
            if isinstance(v, (int, float)) and not isinstance(v, bool) and re.search(r"%|^Share$|^Cumul\. %$", labels[ci]) \
                    and "%" not in cell.number_format and not re.search(r"pp|y\"", cell.number_format):
                E(f"{sheet}!{cell.coordinate}: column '{labels[ci]}' holds shares — store 0–1 values with format 0.0% (not {cell.number_format!r})")
    for sk, rx_t, need in _NEED_COLS[typ]:
        if sk == key and re.search(rx_t, title):
            for rx_c in need:
                if not any(re.search(rx_c, lab, re.I) for lab in labels):
                    W(f"{sheet} / {title}: column matching '{rx_c}' not found — the site's chart for this table needs it (labels: {labels})")
